"""添付ファイルからのテキスト抽出。

.docx / .xlsx / .pdf は書式を捨ててテキストのみ取り出す。
テキスト系ファイルは UTF-8 → Shift_JIS(cp932)の順で読みを試みる。
"""

import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".log",
    ".docx", ".xlsx", ".xlsm", ".pdf",
}


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".docx":
        return _from_docx(data)
    if ext in (".xlsx", ".xlsm"):
        return _from_xlsx(data)
    if ext == ".pdf":
        return _from_pdf(data)
    return _from_plain_text(data)


def _from_plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp932"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _from_docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    lines = []
    # 本文の出現順に段落と表を取り出す
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            lines.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _from_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        if len(wb.worksheets) > 1:
            lines.append(f"■ シート: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append("\t".join(cells).rstrip("\t"))
        lines.append("")
    return "\n".join(lines).strip()


def _from_pdf(data: bytes) -> str:
    from pdfminer.high_level import extract_text as pdf_extract

    text = pdf_extract(io.BytesIO(data))
    # pdfminer は改行が多めに入るので、3連続以上の改行を2つにまとめる
    import re

    return re.sub(r"\n{3,}", "\n\n", text).strip()
